#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador de dados 100% FICTÍCIOS: indústria de materiais de construção
(tubos, conexões e acessórios). Seed fixa => resultado reproduzível.

Nenhum dado real de empresa foi usado. Nomes de clientes e transportadoras
são inventados; qualquer coincidência é acidental.
Os problemas de qualidade (duplicados, nulos, textos inconsistentes etc.)
são injetados de propósito, para treino de limpeza de dados.
"""
import json, sys, datetime as dt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SEED = 42
rng = np.random.default_rng(SEED)
INI, FIM = dt.date(2024, 1, 1), dt.date(2025, 12, 31)
N_LINHAS_BASE = 19700          # + ~300 duplicados injetados = ~20 mil
GAB = {}                       # gabarito: contagem de problemas injetados
OUT_XLSX = sys.argv[1] if len(sys.argv) > 1 else "dados_ficticios.xlsx"
OUT_GAB = sys.argv[2] if len(sys.argv) > 2 else "gabarito_qualidade.json"


def D(n):
    return dt.timedelta(days=int(n))


def pick(n, k, exclude=()):
    pool = np.setdiff1d(np.arange(n), np.array(list(exclude), dtype=int))
    return rng.choice(pool, size=min(k, len(pool)), replace=False)


def setcol(df, idx, col, vals):
    if not hasattr(vals, "__len__") or isinstance(vals, str):
        vals = [vals] * len(idx)
    for i, v in zip(idx, vals):
        df.at[int(i), col] = v


def insert_dups(df, k):
    """Duplica k linhas: metade logo abaixo do original, metade espalhada."""
    n = len(df)
    src = rng.choice(n, size=k, replace=False)
    dup = df.iloc[src].copy()
    keys = np.where(rng.random(k) < 0.5, src + 0.5, rng.uniform(0, n, k))
    base = df.copy()
    base["_k"] = np.arange(n, dtype=float)
    dup["_k"] = keys
    out = pd.concat([base, dup]).sort_values("_k", kind="stable").drop(columns="_k")
    return out.reset_index(drop=True)


# ----------------------------------------------------------------- REGIÕES
UFS = {
    'SC': ('Santa Catarina', 'Sul', .12, ['Joinville', 'Florianópolis', 'Blumenau', 'Chapecó']),
    'PR': ('Paraná', 'Sul', .08, ['Curitiba', 'Londrina', 'Maringá']),
    'RS': ('Rio Grande do Sul', 'Sul', .07, ['Porto Alegre', 'Caxias do Sul', 'Pelotas']),
    'SP': ('São Paulo', 'Sudeste', .22, ['São Paulo', 'Campinas', 'Ribeirão Preto', 'Santos']),
    'MG': ('Minas Gerais', 'Sudeste', .08, ['Belo Horizonte', 'Uberlândia', 'Juiz de Fora']),
    'RJ': ('Rio de Janeiro', 'Sudeste', .06, ['Rio de Janeiro', 'Niterói', 'Campos dos Goytacazes']),
    'ES': ('Espírito Santo', 'Sudeste', .02, ['Vitória', 'Vila Velha']),
    'GO': ('Goiás', 'Centro-Oeste', .05, ['Goiânia', 'Anápolis']),
    'DF': ('Distrito Federal', 'Centro-Oeste', .03, ['Brasília']),
    'MT': ('Mato Grosso', 'Centro-Oeste', .03, ['Cuiabá', 'Rondonópolis']),
    'MS': ('Mato Grosso do Sul', 'Centro-Oeste', .02, ['Campo Grande', 'Dourados']),
    'BA': ('Bahia', 'Nordeste', .05, ['Salvador', 'Feira de Santana']),
    'PE': ('Pernambuco', 'Nordeste', .04, ['Recife', 'Caruaru']),
    'CE': ('Ceará', 'Nordeste', .04, ['Fortaleza', 'Juazeiro do Norte']),
    'MA': ('Maranhão', 'Nordeste', .01, ['São Luís']),
    'RN': ('Rio Grande do Norte', 'Nordeste', .01, ['Natal']),
    'PB': ('Paraíba', 'Nordeste', .01, ['João Pessoa']),
    'AL': ('Alagoas', 'Nordeste', .005, ['Maceió']),
    'SE': ('Sergipe', 'Nordeste', .005, ['Aracaju']),
    'PI': ('Piauí', 'Nordeste', .005, ['Teresina']),
    'PA': ('Pará', 'Norte', .03, ['Belém', 'Santarém']),
    'AM': ('Amazonas', 'Norte', .02, ['Manaus']),
    'TO': ('Tocantins', 'Norte', .01, ['Palmas']),
    'RO': ('Rondônia', 'Norte', .01, ['Porto Velho']),
    'AC': ('Acre', 'Norte', .005, ['Rio Branco']),
    'AP': ('Amapá', 'Norte', .005, ['Macapá']),
    'RR': ('Roraima', 'Norte', .005, ['Boa Vista']),
}
uf_list = list(UFS)
uf_w = np.array([UFS[u][2] for u in uf_list]); uf_w /= uf_w.sum()
regioes = pd.DataFrame({
    'uf': uf_list,
    'nome_estado': [UFS[u][0] for u in uf_list],
    'regiao': [UFS[u][1] for u in uf_list],
})

# ---------------------------------------------------------------- PRODUTOS
prods = []  # (familia, descricao, preco_tabela, unidade)


def add(fam, fmt, items, un='UN', jitter=True):
    for d, p in items:
        p = p * (rng.uniform(.95, 1.05) if jitter else 1)
        prods.append((fam, fmt % d, round(p, 2), un))


add('Tubos', 'Tubo PVC Esgoto %dmm 6m', [(40, 28), (50, 38), (75, 62), (100, 95), (150, 190)], 'BR')
add('Tubos', 'Tubo PVC Soldável %dmm 6m', [(20, 14), (25, 19), (32, 28), (40, 42), (50, 58), (60, 80), (75, 120), (85, 150), (110, 210)], 'BR')
add('Tubos', 'Tubo Eletroduto PVC %dmm 3m', [(20, 9), (25, 12), (32, 17), (40, 24), (50, 33)], 'BR')
add('Tubos', 'Tubo PEAD PE80 %dmm rolo 100m', [(20, 62), (32, 98), (50, 185), (63, 260)], 'RL')
add('Tubos', 'Tubo PPR %dmm 4m', [(20, 22), (25, 30), (32, 46), (40, 70), (50, 105)], 'BR')
add('Conexões', 'Joelho 90° Soldável %dmm', [(20, 1.8), (25, 2.4), (32, 4.2), (40, 7.5), (50, 11), (60, 19), (75, 34)])
add('Conexões', 'Joelho 45° Soldável %dmm', [(20, 2.0), (25, 2.7), (32, 4.6), (40, 8.2), (50, 12)])
add('Conexões', 'Luva Soldável %dmm', [(20, 1.4), (25, 2.0), (32, 3.5), (40, 6.0), (50, 9.0), (60, 15)])
add('Conexões', 'Tê Soldável %dmm', [(20, 2.6), (25, 3.4), (32, 6.2), (40, 10.5), (50, 15)])
add('Conexões', 'Curva Esgoto 90° %dmm', [(40, 6), (50, 9), (75, 18), (100, 31)])
add('Conexões', 'Joelho Esgoto 45° %dmm', [(40, 5), (50, 8), (75, 15), (100, 27), (150, 70)])
add('Conexões', 'Junção Simples Esgoto %s', [('50x50mm', 11), ('100x50mm', 24), ('100x100mm', 38)])
add('Conexões', 'Redução Soldável %s', [('25x20mm', 2.2), ('32x25mm', 3.6), ('40x32mm', 6.3), ('50x40mm', 9.4), ('60x50mm', 15)])
add('Conexões', 'Adaptador Soldável Curto %dmm', [(20, 1.9), (25, 2.6), (32, 4.4)])
add('Conexões', 'União Soldável %dmm', [(20, 7), (25, 10), (32, 16), (40, 29), (50, 42)])
add('Acessórios', 'Registro de Esfera Soldável %dmm', [(25, 24), (32, 38), (50, 85)])
add('Acessórios', '%s', [('Registro de Gaveta 3/4"', 52), ('Ralo Sifonado 100mm', 14), ('Ralo Linear 70cm', 120),
                          ('Sifão Sanfonado Universal', 9), ('Caixa Sifonada 150x150mm', 26),
                          ('Torneira de Boia 3/4"', 28), ('Flexível Inox 40cm', 18),
                          ('Abraçadeira Metálica 3/4"', 2.2), ('Válvula de Retenção Vertical 1"', 45),
                          ('Mangueira de Jardim 1/2" 50m', 85), ('Tampão Esgoto 100mm', 4.5),
                          ('Anel de Vedação Esgoto 100mm', 3.5), ('Caixa de Gordura 250mm', 49),
                          ('Fita Veda Rosca 18mm x 50m', 7.5), ('Fita Veda Rosca 18mm x 10m', 3.2)])
add('Acessórios', 'Caixa d\'Água Polietileno %dL', [(500, 420), (1000, 790)])
add('Acessórios', 'Adesivo para PVC %s', [('75g', 11), ('175g', 19), ('850g', 68)])

produtos = pd.DataFrame(prods, columns=['familia', 'descricao', 'preco_tabela', 'unidade_medida'])
produtos.insert(0, 'sku', [f"PRD-{i + 1:04d}" for i in range(len(produtos))])
NS = len(produtos)

fam_m = {'Tubos': .34, 'Conexões': .42, 'Acessórios': .50}   # margem sobre preço de tabela
m = produtos['familia'].map(fam_m).values + rng.normal(0, .05, NS)
low = rng.choice(np.where(produtos['familia'] == 'Tubos')[0], 5, replace=False)
m[low] = rng.uniform(.17, .23, 5)                              # itens de margem baixa
produtos['custo'] = np.round(produtos['preco_tabela'] * (1 - np.clip(m, .1, .7)), 2)

# peso de demanda por SKU
fam_w = {'Tubos': 1.0, 'Conexões': 1.5, 'Acessórios': 0.8}
w = produtos['familia'].map(fam_w).values * rng.lognormal(0, .7, NS) / produtos['preco_tabela'].values ** .2
rank = np.argsort(-w)
top = rank[:40]
rest = rank[40:]
parados = rng.choice(rest[:-5], 8, replace=False)             # param: sem venda há 120+ dias
resto2 = np.setdiff1d(rest, parados)
lentos = rng.choice(resto2, 6, replace=False)
w[lentos] *= .08
at_risk = rng.choice(top, 10, replace=False)
excesso = rng.choice(np.setdiff1d(resto2, lentos), 8, replace=False)
cutoff = np.full(NS, dt.date(2100, 1, 1), dtype=object)
for i in parados:
    cutoff[i] = INI + D(rng.integers((dt.date(2025, 2, 15) - INI).days, (dt.date(2025, 8, 25) - INI).days))

# ---------------------------------------------------------------- CLIENTES
palavras = ['Aurora', 'Horizonte', 'Planalto', 'Vale Verde', 'Pioneira', 'Atlântica', 'Serra Azul', 'Central',
            'Nova Era', 'Litoral', 'Imperial', 'Bandeirante', 'Solar', 'Alvorada', 'Cruzeiro', 'Progresso',
            'Triunfo', 'Monte Alto', 'Rio Claro', 'Estrela', 'Fênix', 'Vanguarda', 'Esmeralda', 'Colina',
            'Tropical', 'Meridiano', 'Safira', 'Jacarandá', 'Ipê', 'Araucária', 'Cedro', 'Girassol', 'Boreal',
            'Orion', 'Platina', 'Cobalto', 'Lótus', 'Panorama', 'Sertaneja', 'Marajó']
sufixos = ['Materiais', 'Construção', 'Hidráulica', 'Atacado', 'Suprimentos', 'Comercial']
canais = {'Distribuidor': 130, 'Home Center': 70, 'Construtora': 150, 'E-commerce': 40}
nomes_usados, cli_rows = set(), []
for canal, qtd in canais.items():
    c = 0
    while c < qtd:
        a = palavras[rng.integers(len(palavras))]
        b = sufixos[rng.integers(len(sufixos))]
        if canal == 'Distribuidor':
            nome = f"Distribuidora {a} {b} Ltda"
        elif canal == 'Home Center':
            nome = f"{a} Home Center {b.split()[0]}" if rng.random() < .5 else f"Casa & Obra {a} {b}"
        elif canal == 'Construtora':
            nome = f"Construtora {a} {b[:4]}" if rng.random() < .5 else f"{a} Engenharia {b[:3]} S/A"
        else:
            nome = f"{a} {b} Online"
        if nome in nomes_usados:
            continue
        nomes_usados.add(nome)
        uf = rng.choice(uf_list, p=uf_w)
        cidade = UFS[uf][3][rng.integers(len(UFS[uf][3]))]
        cli_rows.append((nome, canal, cidade, uf))
        c += 1
perm = rng.permutation(len(cli_rows))
cli_rows = [cli_rows[i] for i in perm]
clientes = pd.DataFrame(cli_rows, columns=['nome_cliente', 'canal', 'cidade', 'uf'])
clientes.insert(0, 'id_cliente', [f"CLI-{i + 1:04d}" for i in range(len(clientes))])
NC = len(clientes)
reg_map = dict(zip(regioes.uf, regioes.regiao))
clientes['regiao'] = clientes['uf'].map(reg_map)
canal_freq = {'Distribuidor': 1.6, 'Home Center': 2.0, 'Construtora': .8, 'E-commerce': 1.2}
pc = rng.lognormal(0, .9, NC) * clientes['canal'].map(canal_freq).values
pc /= pc.sum()

# ------------------------------------------------------------------ VENDAS
dias = [INI + D(i) for i in range((FIM - INI).days + 1)]
sazon = [.85, .9, 1.0, 1.0, 1.0, .95, 1.0, 1.05, 1.1, 1.1, 1.1, .95]
pw = np.array([sazon[d.month - 1] * (1.08 if d.year == 2025 else 1.0) *
               (1.0 if d.weekday() < 5 else (.35 if d.weekday() == 5 else .02)) for d in dias])
pw /= pw.sum()
n_ped_est = 7200
datas = sorted(rng.choice(len(dias), size=n_ped_est, p=pw))
cli_ped = rng.choice(NC, size=n_ped_est, p=pc)
lam = {'Distribuidor': 2.0, 'Home Center': 2.4, 'Construtora': 2.2, 'E-commerce': .4}
q_mu = {'Tubos': np.log(10), 'Conexões': np.log(30), 'Acessórios': np.log(6)}
ch_mult = {'Distribuidor': 3.0, 'Construtora': 1.5, 'Home Center': 2.0, 'E-commerce': .4}
ch_desc = {'Distribuidor': .12, 'Construtora': .08, 'Home Center': .10, 'E-commerce': .03}
pt = produtos['preco_tabela'].values
sku_scale = np.where(pt > 200, .15, np.where(pt > 80, .4, np.where(pt > 30, .7, 1.0)))
fam_arr = produtos['familia'].values
sku_arr = produtos['sku'].values
tot_dias = (FIM - INI).days

linhas, ped_n = [], 0
for di, ci in zip(datas, cli_ped):
    d = dias[di]
    cl = clientes.iloc[ci]
    ped_n += 1
    k = int(np.clip(1 + rng.poisson(lam[cl.canal]), 1, 8))
    ativo = np.array([c > d for c in cutoff], dtype=float)
    pr = w * ativo
    pr /= pr.sum()
    skus = rng.choice(NS, size=k, replace=False, p=pr)
    idx_preco = .94 + .06 * (di / tot_dias)
    for j, s in enumerate(skus, 1):
        q = max(1, int(round(rng.lognormal(q_mu[fam_arr[s]], .7) * ch_mult[cl.canal] * sku_scale[s])))
        desc = float(np.clip(ch_desc[cl.canal] + rng.normal(0, .03), 0, .25))
        linhas.append((f"PED-{ped_n:06d}", j, d, cl.id_cliente, sku_arr[s], q,
                       round(pt[s] * idx_preco, 2), round(desc, 2)))
keep = np.sort(rng.choice(len(linhas), N_LINHAS_BASE, replace=False))
linhas = [linhas[i] for i in keep]
vendas_clean = pd.DataFrame(linhas, columns=['id_pedido', 'item', 'data_pedido', 'id_cliente', 'sku',
                                             'quantidade', 'preco_unitario', 'desconto_pct'])
vendas_clean['item'] = vendas_clean.groupby('id_pedido').cumcount() + 1
vendas_clean['canal'] = vendas_clean['id_cliente'].map(clientes.set_index('id_cliente')['canal'])

# --------------------------------------------------------------- ENTREGAS
ped = (vendas_clean.groupby('id_pedido')
       .agg(data=('data_pedido', 'first'), cli=('id_cliente', 'first'), qtd=('quantidade', 'sum'))
       .reset_index())
cli_reg = clientes.set_index('id_cliente')['regiao']
base_t = {'Sul': 2, 'Sudeste': 3, 'Centro-Oeste': 5, 'Nordeste': 6, 'Norte': 9}
transp_reg = {
    'Sul': {'TransLog Sul': .55, 'Via Brasil Cargas': .30, 'EcoTrans': .15},
    'Sudeste': {'TransLog Sul': .20, 'Via Brasil Cargas': .30, 'EcoTrans': .35, 'Expresso Centro': .15},
    'Centro-Oeste': {'Via Brasil Cargas': .30, 'EcoTrans': .20, 'Expresso Centro': .50},
    'Nordeste': {'Via Brasil Cargas': .30, 'Rápido Norte': .40, 'Expresso Centro': .30},
    'Norte': {'Rápido Norte': .60, 'Via Brasil Cargas': .40},
}
shift = {'TransLog Sul': -.2, 'Via Brasil Cargas': .3, 'EcoTrans': 0.0, 'Expresso Centro': .9, 'Rápido Norte': 1.2}
short = {'TransLog Sul': .04, 'Via Brasil Cargas': .07, 'EcoTrans': .12, 'Expresso Centro': .09, 'Rápido Norte': .10}
BUF = 3
ent = []
for r in ped.itertuples():
    reg = cli_reg[r.cli]
    tr = rng.choice(list(transp_reg[reg]), p=list(transp_reg[reg].values()))
    bt = base_t[reg]
    exped = int(rng.integers(1, 4))
    trans = bt + shift[tr] + rng.normal(0, 1.3)
    if rng.random() < .06:
        trans += rng.uniform(2, 6)
    trans = max(1, int(round(trans)))
    d_exp = r.data + D(exped)
    d_prom = r.data + D(exped + bt + BUF)
    d_ent = r.data + D(exped + trans)
    if rng.random() < short[tr]:
        qe = int(np.clip(int(r.qtd * rng.uniform(.5, .97)), 0, max(r.qtd - 1, 0)))
    else:
        qe = int(r.qtd)
    if d_ent > FIM:
        ent.append((r.id_pedido, tr, d_exp if d_exp <= FIM else None, d_prom, None, int(r.qtd), None, 'Em trânsito'))
    else:
        ent.append((r.id_pedido, tr, d_exp, d_prom, d_ent, int(r.qtd), qe, 'Entregue'))
entregas_clean = pd.DataFrame(ent, columns=['id_pedido', 'transportadora', 'data_expedicao', 'data_prometida',
                                            'data_entrega', 'qtd_pedida', 'qtd_entregue', 'status'])

# ----------------------------------------------------------------- ESTOQUE
vc = vendas_clean.copy()
vc['d'] = pd.to_datetime(vc['data_pedido'])
ult180 = vc[vc['d'] >= pd.Timestamp(FIM - D(179))].groupby('sku')['quantidade'].sum() / 180
tot = vc.groupby('sku')['quantidade'].sum()
last = vc.groupby('sku')['d'].max()
est = []
for i, s in enumerate(sku_arr):
    dd = ult180.get(s, 0.0)
    if i in parados:
        dd = tot.get(s, 1) / max((cutoff[i] - INI).days, 1)
        cov = rng.uniform(120, 260)
    elif dd <= 0:
        dd = tot.get(s, 1) / 730
        cov = rng.uniform(60, 150)
    elif i in at_risk:
        cov = rng.uniform(2, 12)
    elif i in excesso:
        cov = rng.uniform(90, 200)
    else:
        cov = float(np.clip(rng.lognormal(np.log(28), .35), 12, 65))
    qtd = int(round(dd * cov))
    minimo = max(int(round(dd * 15)), 5)
    ult = last[s].date() if s in last.index else FIM - D(rng.integers(150, 400))
    est.append((s, qtd, minimo, ult, FIM))
estoque_clean = pd.DataFrame(est, columns=['sku', 'qtd_atual', 'estoque_minimo', 'ultima_movimentacao', 'data_posicao'])

# =============================================================== SUJEIRA
# ---- vendas
v = vendas_clean.drop(columns='canal').astype(object).copy()
n = len(v)
g = {}
i1 = pick(n, int(.03 * n)); setcol(v, i1, 'data_pedido', [v.at[int(i), 'data_pedido'].strftime('%d/%m/%Y') for i in i1])
i2 = pick(n, int(.01 * n), i1); setcol(v, i2, 'data_pedido', [v.at[int(i), 'data_pedido'].strftime('%Y-%m-%d') for i in i2])
g['data_pedido_como_texto_dd/mm/aaaa'] = len(i1); g['data_pedido_como_texto_aaaa-mm-dd'] = len(i2)
i = pick(n, int(.004 * n)); setcol(v, i, 'id_cliente', None); g['id_cliente_nulo'] = len(i)
i = pick(n, int(.003 * n)); setcol(v, i, 'id_cliente', [f"CLI-{int(x):04d}" for x in rng.integers(900, 999, len(i))]); g['id_cliente_orfao'] = len(i)
i = pick(n, int(.02 * n)); setcol(v, i, 'sku', [v.at[int(j), 'sku'] + ' ' for j in i]); g['sku_espaco_final'] = len(i)
i = pick(n, int(.015 * n)); setcol(v, i, 'sku', [str(v.at[int(j), 'sku']).strip().lower() for j in i]); g['sku_minusculo'] = len(i)
i = pick(n, int(.01 * n)); setcol(v, i, 'sku', [str(v.at[int(j), 'sku']).strip().replace('-', '') for j in i]); g['sku_sem_hifen'] = len(i)
i = pick(n, int(.003 * n)); setcol(v, i, 'quantidade', [-int(v.at[int(j), 'quantidade']) for j in i]); g['quantidade_negativa'] = len(i)
i = pick(n, 12); setcol(v, i, 'quantidade', [int(v.at[int(j), 'quantidade']) * 100 for j in i]); g['quantidade_outlier_x100'] = len(i)
i = pick(n, 15); setcol(v, i, 'quantidade', None); g['quantidade_nula'] = len(i)
i = pick(n, int(.01 * n)); setcol(v, i, 'preco_unitario', [f"{v.at[int(j), 'preco_unitario']:.2f}".replace('.', ',') for j in i]); g['preco_como_texto_virgula'] = len(i)
i = pick(n, int(.005 * n)); setcol(v, i, 'preco_unitario', None); g['preco_nulo'] = len(i)
i = pick(n, int(.02 * n)); setcol(v, i, 'desconto_pct', None); g['desconto_nulo'] = len(i)
i = [j for j in pick(n, int(.012 * n)) if v.at[int(j), 'desconto_pct'] is not None]; setcol(v, i, 'desconto_pct', [round(float(v.at[int(j), 'desconto_pct']) * 100) for j in i]); g['desconto_escala_errada_(10_em_vez_de_0,10)'] = len(i)
v = insert_dups(v, 300); g['vendas_linhas_duplicadas_exatas'] = 300
vendas = v
GAB['vendas'] = g

# ---- clientes
c = clientes.drop(columns='regiao').astype(object).copy()
nc = len(c); g = {}
nome_uf = {u: UFS[u][0] for u in UFS}
i = pick(nc, int(.08 * nc)); setcol(c, i, 'uf', [c.at[int(j), 'uf'].lower() for j in i]); g['uf_minuscula'] = len(i)
i = pick(nc, int(.04 * nc)); setcol(c, i, 'uf', [nome_uf[c.at[int(j), 'uf'].upper()] for j in i]); g['uf_nome_por_extenso'] = len(i)
i = pick(nc, int(.03 * nc)); setcol(c, i, 'uf', [' ' + str(c.at[int(j), 'uf']) + ' ' for j in i]); g['uf_com_espacos'] = len(i)
i = pick(nc, 4); setcol(c, i, 'uf', None); g['uf_nula'] = len(i)
var_canal = {'Home Center': ['home center', 'HomeCenter', 'Home-Center'], 'E-commerce': ['Ecommerce', 'e-commerce', 'E-Commerce '],
             'Distribuidor': ['Distribuidora', 'distribuidor'], 'Construtora': ['construtora', 'Construtora ']}
i = pick(nc, int(.12 * nc)); setcol(c, i, 'canal', [var_canal[c.at[int(j), 'canal']][rng.integers(len(var_canal[c.at[int(j), 'canal']]))] for j in i]); g['canal_variacoes_de_grafia'] = len(i)
i = pick(nc, 2); setcol(c, i, 'canal', None); g['canal_nulo'] = len(i)
i = pick(nc, int(.03 * nc)); setcol(c, i, 'cidade', None); g['cidade_nula'] = len(i)
i = pick(nc, int(.05 * nc)); setcol(c, i, 'cidade', [str(c.at[int(j), 'cidade']).upper() for j in i if c.at[int(j), 'cidade']]) if False else None
j_ok = [j for j in pick(nc, int(.05 * nc)) if c.at[int(j), 'cidade'] is not None]
setcol(c, j_ok, 'cidade', [str(c.at[int(j), 'cidade']).upper() for j in j_ok]); g['cidade_maiuscula'] = len(j_ok)
dups = []
base_ids = rng.choice(nc, 12, replace=False)
for k, j in enumerate(base_ids):
    r = c.iloc[int(j)].copy()
    r['id_cliente'] = f"CLI-{nc + 1 + k:04d}"
    r['nome_cliente'] = str(r['nome_cliente']).replace('Ltda', 'LTDA').replace('S/A', 'SA') + (' ' if k % 2 else '')
    dups.append(r)
c = pd.concat([c, pd.DataFrame(dups)], ignore_index=True)
g['clientes_duplicados_com_id_novo'] = 12
clientes_sujo = c
GAB['clientes'] = g

# ---- produtos
p = produtos.astype(object).copy()
npd = len(p); g = {}
var_fam = {'Conexões': ['Conexoes', 'CONEXÕES', 'conexões'], 'Tubos': ['tubos', 'TUBOS', 'Tubo'], 'Acessórios': ['Acessorios', 'acessórios', 'ACESSÓRIOS']}
i = pick(npd, int(.15 * npd)); setcol(p, i, 'familia', [var_fam[p.at[int(j), 'familia']][rng.integers(3)] for j in i]); g['familia_variacoes_de_grafia'] = len(i)
i = pick(npd, 6); setcol(p, i, 'preco_tabela', ["R$ " + f"{p.at[int(j), 'preco_tabela']:.2f}".replace('.', ',') for j in i]); g['preco_tabela_como_texto_R$'] = len(i)
i = pick(npd, 3); setcol(p, i, 'custo', None); g['custo_nulo'] = len(i)
ok = [j for j in pick(npd, 12) if isinstance(p.at[int(j), 'preco_tabela'], float)][:2]
setcol(p, ok, 'custo', [round(p.at[int(j), 'preco_tabela'] * 1.3, 2) for j in ok]); g['custo_maior_que_preco'] = len(ok)
i = pick(npd, 8); setcol(p, i, 'descricao', [str(p.at[int(j), 'descricao']) + '  ' for j in i]); g['descricao_espacos_extras'] = len(i)
i = pick(npd, 3); dp = p.iloc[i].copy()
dp['preco_tabela'] = [("R$ 0,00" if False else x) for x in dp['preco_tabela']]
p = pd.concat([p, dp], ignore_index=True); g['sku_duplicado_na_tabela'] = 3
produtos_sujo = p
GAB['produtos'] = g

# ---- entregas
e = entregas_clean.astype(object).copy()
ne = len(e); g = {}
var_tr = {'TransLog Sul': ['Translog Sul', 'TRANSLOG SUL', 'Trans Log Sul', 'TransLog Sul '],
          'Rápido Norte': ['Rapido Norte', 'RÁPIDO NORTE', 'Rápido  Norte'],
          'Via Brasil Cargas': ['Via Brasil', 'via brasil cargas', 'Via Brasil Cargas '],
          'EcoTrans': ['Ecotrans', 'ECOTRANS', 'Eco Trans'],
          'Expresso Centro': ['Expresso Centro ', 'EXPRESSO CENTRO', 'Exp. Centro']}
i = pick(ne, int(.10 * ne)); setcol(e, i, 'transportadora', [var_tr[e.at[int(j), 'transportadora']][rng.integers(len(var_tr[e.at[int(j), 'transportadora']]))] for j in i]); g['transportadora_variacoes_de_grafia'] = len(i)
i = pick(ne, int(.03 * ne)); setcol(e, i, 'data_prometida', [e.at[int(j), 'data_prometida'].strftime('%d/%m/%Y') for j in i]); g['data_prometida_como_texto'] = len(i)
ok = [j for j in pick(ne, int(.05 * ne)) if e.at[int(j), 'data_entrega'] is not None][:int(.03 * ne)]
setcol(e, ok, 'data_entrega', [e.at[int(j), 'data_entrega'].strftime('%d/%m/%Y') for j in ok]); g['data_entrega_como_texto'] = len(ok)
ent_idx = [j for j in range(ne) if e.at[j, 'status'] == 'Entregue']
okn = list(rng.choice(ent_idx, int(.008 * len(ent_idx)), replace=False))
setcol(e, okn, 'data_entrega', None); g['data_entrega_nula_em_pedido_entregue'] = len(okn)
rest_e = [j for j in ent_idx if j not in set(okn) and isinstance(e.at[j, 'data_entrega'], dt.date)]
okb = list(rng.choice(rest_e, int(.003 * ne), replace=False))
setcol(e, okb, 'data_entrega', [e.at[int(j), 'data_expedicao'] - D(rng.integers(1, 4)) for j in okb]); g['data_entrega_antes_da_expedicao'] = len(okb)
rest_e = [j for j in ent_idx if j not in set(okn) | set(okb)]
oko = list(rng.choice(rest_e, int(.005 * ne), replace=False))
setcol(e, oko, 'qtd_entregue', [int(e.at[int(j), 'qtd_pedida'] * rng.uniform(1.02, 1.1)) + 1 for j in oko]); g['qtd_entregue_maior_que_pedida'] = len(oko)
rest_e = [j for j in ent_idx if j not in set(okn) | set(okb) | set(oko)]
okq = list(rng.choice(rest_e, int(.003 * ne), replace=False))
setcol(e, okq, 'qtd_entregue', None); g['qtd_entregue_nula'] = len(okq)
i = pick(ne, int(.05 * ne)); setcol(e, i, 'status', [str(e.at[int(j), 'status']).lower() if k % 2 else str(e.at[int(j), 'status']).upper() for k, j in enumerate(i)]); g['status_caixa_inconsistente'] = len(i)
drop = pick(ne, 25); e = e.drop(index=drop).reset_index(drop=True); g['pedidos_sem_registro_de_entrega'] = 25
orf = []
for k in range(10):
    r = e.iloc[int(rng.integers(len(e)))].copy(); r['id_pedido'] = f"PED-{99000 + k}"; orf.append(r)
e = pd.concat([e, pd.DataFrame(orf)], ignore_index=True); g['entrega_de_pedido_inexistente_em_vendas'] = 10
e = insert_dups(e, 80); g['entregas_linhas_duplicadas'] = 80
entregas = e
GAB['entregas'] = g

# ---- estoque
s_ = estoque_clean.astype(object).copy()
ns = len(s_); g = {}
i = pick(ns, 3); setcol(s_, i, 'qtd_atual', [-abs(int(s_.at[int(j), 'qtd_atual'])) - 1 for j in i]); g['qtd_atual_negativa'] = len(i)
i = pick(ns, 4); setcol(s_, i, 'ultima_movimentacao', None); g['ultima_movimentacao_nula'] = len(i)
i = pick(ns, 3, i); setcol(s_, i, 'ultima_movimentacao', 'N/A'); g['ultima_movimentacao_texto_N/A'] = len(i)
i = pick(ns, 3); setcol(s_, i, 'estoque_minimo', None); g['estoque_minimo_nulo'] = len(i)
i = pick(ns, 2, i); setcol(s_, i, 'estoque_minimo', 0); g['estoque_minimo_zero'] = len(i)
i = pick(ns, 5); setcol(s_, i, 'sku', [s_.at[int(j), 'sku'].lower() + ' ' for j in i]); g['sku_minusculo_ou_com_espaco'] = len(i)
orf = []
for k in range(2):
    r = s_.iloc[int(rng.integers(ns))].copy(); r['sku'] = f"PRD-{990 + k:04d}"; orf.append(r)
s_ = pd.concat([s_, pd.DataFrame(orf)], ignore_index=True); g['sku_inexistente_em_produtos'] = 2
s_ = insert_dups(s_, 4); g['estoque_linhas_duplicadas'] = 4
estoque = s_
GAB['estoque'] = g

# ================================================================ EXCEL
FONT = 'Arial'
HFILL = PatternFill('solid', start_color='1F3864')


def conv(x):
    if x is None or x is pd.NaT:
        return None
    if isinstance(x, (float, np.floating)):
        if np.isnan(x):
            return None
        x = float(x)
        return int(x) if x.is_integer() else x
    if isinstance(x, np.integer):
        return int(x)
    if isinstance(x, pd.Timestamp):
        return x.date()
    return x


def write_df(wb, name, df, widths=None, pct_cols=(), money_cols=()):
    ws = wb.create_sheet(name)
    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=10)
        cell.fill = HFILL
        cell.alignment = Alignment(horizontal='center')
    cols = list(df.columns)
    for row in df.itertuples(index=False):
        ws.append([conv(x) for x in row])
    for r in ws.iter_rows(min_row=2):
        for cell in r:
            cell.font = Font(name=FONT, size=10)
            if isinstance(cell.value, dt.date):
                cell.number_format = 'dd/mm/yyyy'
            elif cols[cell.column - 1] in pct_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '0%'
            elif cols[cell.column - 1] in money_cols and isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'
    for k, col in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(k)].width = (widths or {}).get(col, max(14, len(col) + 4))
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


wb = Workbook()
ws0 = wb.active
ws0.title = 'LEIA-ME'
linhas_leia = [
    ("DADOS 100% FICTÍCIOS — PROJETO DE PORTFÓLIO (EXCEL + POWER BI)", 'h'),
    ("Indústria fictícia de materiais de construção (tubos, conexões e acessórios).", ''),
    ("Nenhuma informação real de empresa foi usada ou simulada. Nomes de clientes e transportadoras são inventados;", ''),
    ("qualquer coincidência é acidental. Este arquivo não representa nenhuma marca.", ''),
    ("Gerado por script em Python com seed fixa (reprodutível). Período: 01/01/2024 a 31/12/2025.", ''),
    ("", ''),
    ("TABELAS", 'h'),
    ("vendas: itens de pedido (uma linha por item). Faturamento = quantidade × preço unitário × (1 − desconto).", ''),
    ("clientes: cadastro com canal, cidade e UF.", ''),
    ("produtos: SKU, família, descrição, unidade, custo unitário e preço de tabela (2025).", ''),
    ("regioes: tabela de referência UF → região.", ''),
    ("entregas: uma linha por pedido (transportadora, datas, quantidade pedida × entregue, status).", ''),
    ("estoque: posição de estoque por SKU em 31/12/2025.", ''),
    ("parametros: metas e premissas fictícias usadas no dashboard.", ''),
    ("", ''),
    ("PREMISSAS DO MODELO", 'h'),
    ("Custo unitário é único por SKU (não varia no tempo); o preço de venda sobe gradualmente de 2024 para 2025.", ''),
    ("Pedidos com data de entrega após 31/12/2025 aparecem como 'Em trânsito' (sem data de entrega, de forma legítima).", ''),
    ("", ''),
    ("QUALIDADE DOS DADOS", 'h'),
    ("As tabelas contêm problemas propositais. Categorias: duplicados, nulos, textos inconsistentes (grafia, caixa,", ''),
    ("espaços), tipos misturados (datas e números como texto), valores negativos, outliers, chaves órfãs,", ''),
    ("registros ausentes e datas incoerentes. Localizar e tratar cada um faz parte do exercício.", ''),
]
for t, kind in linhas_leia:
    ws0.append([t])
    ws0.cell(ws0.max_row, 1).font = Font(name=FONT, size=11 if kind == 'h' else 10, bold=(kind == 'h'),
                                         color='C00000' if (kind == 'h' and ws0.max_row == 1) else '000000')
ws0.column_dimensions['A'].width = 120

par = [('Indicador / parâmetro', 'Valor', 'Observação'),
       ('Meta OTIF', 0.92, 'Fictícia. Pedidos no prazo e completos / total'),
       ('Meta lead time médio (dias)', 7, 'Fictícia. Data de entrega − data do pedido'),
       ('Margem % mínima', 0.28, 'Fictícia'),
       ('Cobertura mínima (dias)', 15, 'Abaixo disso = risco de ruptura'),
       ('Cobertura máxima (dias)', 45, 'Acima disso = excesso de estoque'),
       ('Janela de item parado (dias)', 90, 'Sem movimentação há mais de N dias'),
       ('Data de posição do estoque', FIM, 'Foto do estoque')]
ws1 = wb.create_sheet('parametros')
for r in par:
    ws1.append(list(r))
for cell in ws1[1]:
    cell.font = Font(name=FONT, bold=True, color='FFFFFF', size=10); cell.fill = HFILL
for r in ws1.iter_rows(min_row=2):
    for cell in r:
        cell.font = Font(name=FONT, size=10)
ws1['B2'].number_format = '0%'; ws1['B4'].number_format = '0%'; ws1['B8'].number_format = 'dd/mm/yyyy'
ws1.column_dimensions['A'].width = 32; ws1.column_dimensions['B'].width = 14; ws1.column_dimensions['C'].width = 52

write_df(wb, 'vendas', vendas, {'id_pedido': 13, 'data_pedido': 13, 'sku': 12}, pct_cols=('desconto_pct',), money_cols=('preco_unitario',))
write_df(wb, 'clientes', clientes_sujo, {'nome_cliente': 38, 'cidade': 24})
write_df(wb, 'produtos', produtos_sujo, {'descricao': 42}, money_cols=('custo', 'preco_tabela'))
write_df(wb, 'regioes', regioes, {'nome_estado': 24})
write_df(wb, 'entregas', entregas, {'transportadora': 22, 'data_expedicao': 16, 'data_prometida': 16, 'data_entrega': 16})
write_df(wb, 'estoque', estoque, {'ultima_movimentacao': 22, 'data_posicao': 14})
wb.save(OUT_XLSX)

# --------------------------------------------------------------- RESUMO
vc = vendas_clean.merge(produtos[['sku', 'custo', 'familia']], on='sku')
vc['fat'] = vc.quantidade * vc.preco_unitario * (1 - vc.desconto_pct)
vc['cst'] = vc.quantidade * vc.custo
vc['ano'] = pd.to_datetime(vc.data_pedido).dt.year
ec = entregas_clean[entregas_clean.status == 'Entregue']
ontime = pd.to_datetime(ec.data_entrega) <= pd.to_datetime(ec.data_prometida)
full = ec.qtd_entregue >= ec.qtd_pedida
lt = (pd.to_datetime(ec.data_entrega) - ped.set_index('id_pedido').loc[ec.id_pedido, 'data'].pipe(pd.to_datetime).values).dt.days
cob = estoque_clean.set_index('sku')['qtd_atual'] / (ult180.reindex(estoque_clean.sku).fillna(0).values / 1.0).clip(min=1e-9)
GAB['_resumo_base_limpa'] = {
    'linhas_vendas_final': len(vendas), 'pedidos': int(ped.shape[0]),
    'faturamento_total': round(float(vc.fat.sum()), 2),
    'faturamento_2024': round(float(vc[vc.ano == 2024].fat.sum()), 2),
    'faturamento_2025': round(float(vc[vc.ano == 2025].fat.sum()), 2),
    'margem_pct': round(float(1 - vc.cst.sum() / vc.fat.sum()), 4),
    'margem_por_familia': {f: round(float(1 - g_.cst.sum() / g_.fat.sum()), 4) for f, g_ in vc.groupby('familia')},
    'otif': round(float((ontime & full).mean()), 4), 'pct_no_prazo': round(float(ontime.mean()), 4),
    'pct_completo': round(float(full.mean()), 4), 'lead_time_medio': round(float(lt.mean()), 2),
    'skus_parados_90d_ou_mais': int(sum(1 for s in sku_arr if (s not in last.index) or (FIM - last[s].date()).days > 90)),
    'skus_cobertura_<15d': int((cob < 15).sum()), 'skus_cobertura_>45d': int((cob > 45).sum()),
}
with open(OUT_GAB, 'w', encoding='utf-8') as f:
    json.dump(GAB, f, ensure_ascii=False, indent=2, default=str)
print(json.dumps(GAB['_resumo_base_limpa'], ensure_ascii=False, indent=2))
print({k: len(x) for k, x in dict(vendas=vendas, clientes=clientes_sujo, produtos=produtos_sujo, regioes=regioes, entregas=entregas, estoque=estoque).items()})
